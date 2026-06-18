import base64
import hmac
import io
import ipaddress
import logging
import os
import re
import time
import urllib.parse

import fitz  # PyMuPDF
import requests as http_requests
from fastapi import FastAPI, HTTPException, Request

logger = logging.getLogger("esign-api")
logging.basicConfig(level=logging.INFO)

API_KEY = os.environ.get("ESIGN_API_KEY", "")

MAX_PDF_BASE64_BYTES = 70_000_000
MAX_IMAGE_BASE64_BYTES = 14_000_000
MAX_OVERLAY_COUNT = 500
MAX_PDF_URL_COUNT = 20
MAX_FORM_FIELDS = 500
MAX_RENDER_PAGES = 50
MAX_RENDER_DPI = 300

ALLOWED_DOMAIN_SUFFIXES = (
    ".content.force.com",
    ".salesforce.com",
    ".force.com",
    ".documentforce.com",
)

app = FastAPI(title="eSign API", version="1.0.0", docs_url=None, redoc_url=None, openapi_url=None)


def _verify_api_key(request: Request):
    if not API_KEY:
        return
    provided = request.headers.get("X-API-Key", "") or request.headers.get("Esign-Api-Key", "")
    if not provided:
        # Diagnostic: which headers arrived (names only for sensitive ones —
        # never log auth/secret/key values, even truncated).
        _SENSITIVE = ("key", "secret", "authorization", "token", "x-ms-")
        safe_headers = {
            k: ("[redacted]" if any(s in k.lower() for s in _SENSITIVE)
                else (v[:20] + '...' if len(v) > 20 else v))
            for k, v in request.headers.items()
        }
        logger.warning("API key not found in headers. Received headers: %s", safe_headers)
    if not provided or not hmac.compare_digest(provided, API_KEY):
        raise HTTPException(status_code=401, detail="Invalid or missing API key")


def _validate_pdf_url(url: str) -> str:
    if not url or not isinstance(url, str):
        return "Missing or invalid URL"
    parsed = urllib.parse.urlparse(url)
    if parsed.scheme != "https":
        return f"URL scheme must be https, got: {parsed.scheme}"
    hostname = parsed.hostname or ""
    try:
        ip = ipaddress.ip_address(hostname)
        if ip.is_private or ip.is_loopback or ip.is_link_local or ip.is_reserved:
            return "URL points to internal/private address"
    except ValueError:
        pass
    if not any(hostname.endswith(s) for s in ALLOWED_DOMAIN_SUFFIXES):
        return f"URL hostname not in allowed domains: {hostname}"
    return ""


@app.post("/compose-pdf")
async def compose_pdf(request: Request):
    _verify_api_key(request)
    body = await request.json()
    start = time.time()

    pdf_url = body.get("pdf_url")
    overlays = body.get("overlays", [])

    if not pdf_url:
        raise HTTPException(status_code=400, detail="Missing pdf_url")

    url_err = _validate_pdf_url(pdf_url)
    if url_err:
        raise HTTPException(status_code=400, detail=url_err)

    if len(overlays) > MAX_OVERLAY_COUNT:
        raise HTTPException(status_code=400, detail=f"Too many overlays ({len(overlays)}, max {MAX_OVERLAY_COUNT})")

    resp = http_requests.get(pdf_url, timeout=60)
    if resp.status_code != 200:
        raise HTTPException(status_code=502, detail=f"Failed to download PDF: {resp.status_code}")

    doc = fitz.open(stream=resp.content, filetype="pdf")

    for overlay in overlays:
        page_num = overlay.get("page", 1) - 1
        if page_num < 0 or page_num >= len(doc):
            continue

        page = doc[page_num]
        page_rect = page.rect
        pw = page_rect.width
        ph = page_rect.height

        x = (overlay.get("x_pct", 0) / 100.0) * pw
        y = (overlay.get("y_pct", 0) / 100.0) * ph

        overlay_type = overlay.get("type", "")

        if overlay_type == "image":
            image_b64 = overlay.get("image_base64", "")
            if not image_b64:
                continue
            if len(image_b64) > MAX_IMAGE_BASE64_BYTES:
                continue
            img_bytes = base64.b64decode(image_b64)
            w = overlay.get("width", 200)
            h = overlay.get("height", 60)
            rect = fitz.Rect(x, y, x + w, y + h)
            page.insert_image(rect, stream=img_bytes)

        elif overlay_type == "text":
            value = overlay.get("value", "")
            font_size = overlay.get("font_size", 11)
            page.insert_text(
                fitz.Point(x, y),
                value,
                fontsize=font_size,
                fontname="helv",
                color=(0, 0, 0),
            )

        elif overlay_type == "checkbox":
            checked = overlay.get("checked", False)
            if checked:
                size = 12
                rect = fitz.Rect(x, y, x + size, y + size)
                page.draw_rect(rect, color=(0, 0, 0), width=0.5)
                page.draw_line(
                    fitz.Point(x + 2, y + size * 0.55),
                    fitz.Point(x + size * 0.4, y + size - 2),
                    color=(0, 0, 0), width=1.5,
                )
                page.draw_line(
                    fitz.Point(x + size * 0.4, y + size - 2),
                    fitz.Point(x + size - 2, y + 2),
                    color=(0, 0, 0), width=1.5,
                )

    output = io.BytesIO()
    doc.save(output)
    doc.close()

    pdf_b64 = base64.b64encode(output.getvalue()).decode("utf-8")
    logger.info("compose-pdf: overlays=%d duration=%.2fs", len(overlays), time.time() - start)
    return {"pdf_base64": pdf_b64, "success": True}


@app.post("/merge-pdfs")
async def merge_pdfs(request: Request):
    _verify_api_key(request)
    body = await request.json()
    start = time.time()

    pdf_urls = body.get("pdf_urls", [])

    if not pdf_urls or len(pdf_urls) < 2:
        raise HTTPException(status_code=400, detail="At least 2 PDF URLs are required")

    if len(pdf_urls) > MAX_PDF_URL_COUNT:
        raise HTTPException(status_code=400, detail=f"Too many URLs ({len(pdf_urls)}, max {MAX_PDF_URL_COUNT})")

    for url in pdf_urls:
        url_err = _validate_pdf_url(url)
        if url_err:
            raise HTTPException(status_code=400, detail=f"{url_err}: {url}")

    merged = fitz.open()
    page_counts = []

    for i, url in enumerate(pdf_urls):
        resp = http_requests.get(url, timeout=60)
        if resp.status_code != 200:
            merged.close()
            raise HTTPException(status_code=502, detail=f"Failed to download PDF {i + 1}: {resp.status_code}")

        doc = fitz.open(stream=resp.content, filetype="pdf")
        page_counts.append(len(doc))
        merged.insert_pdf(doc)
        doc.close()

    output = io.BytesIO()
    merged.save(output)
    merged.close()

    pdf_b64 = base64.b64encode(output.getvalue()).decode("utf-8")
    total = sum(page_counts)
    logger.info("merge-pdfs: docs=%d pages=%d duration=%.2fs", len(pdf_urls), total, time.time() - start)
    return {
        "pdf_base64": pdf_b64,
        "page_counts": page_counts,
        "total_pages": total,
        "success": True,
    }


@app.post("/extract-pdf-tags")
async def extract_pdf_tags(request: Request):
    _verify_api_key(request)
    body = await request.json()
    start = time.time()

    pdf_b64 = body.get("pdf_base64")
    if not pdf_b64:
        raise HTTPException(status_code=400, detail="Missing pdf_base64")

    if len(pdf_b64) > MAX_PDF_BASE64_BYTES:
        raise HTTPException(status_code=400, detail=f"PDF too large ({len(pdf_b64)} bytes, max {MAX_PDF_BASE64_BYTES})")

    SIGNING_REGEX = re.compile(
        r'\[\s*\[\s*(S|I|N|D|T|CO|EML)(\d+)(?:\.((?:[fsrdnmpxWHOY\-]|\d)+))?\s*\]\s*\]', re.IGNORECASE
    )
    TEXT_TAB_REGEX = re.compile(
        r'\[\s*\[\s*([A-Za-z_]\w*(?:__[rc])?)(?:\.((?:[fsrdnmpxWHOY\-]|\d)+))?\s*\]\s*\]', re.IGNORECASE
    )
    DS_SIGNING_REGEX = re.compile(r'\\([sidnt])(\d+)\\', re.IGNORECASE)
    DS_TEXT_TAB_REGEX = re.compile(r'\\([a-zA-Z][a-zA-Z0-9]*)_(\d+)_(text|number|date|checkbox)\\', re.IGNORECASE)

    SIGNING_TYPE_MAP = {
        'S': 'Signature', 'I': 'Initial', 'N': 'Name', 'D': 'Date',
        'T': 'Title', 'CO': 'Company', 'EML': 'Email'
    }
    DS_SIGNING_TYPE_MAP = {
        's': 'Signature', 'i': 'Initial', 'd': 'Date', 'n': 'Name', 't': 'Title'
    }
    DS_TEXT_TYPE_MAP = {
        'text': 'Text', 'number': 'Number', 'date': 'Date Input', 'checkbox': 'Checkbox'
    }

    SIZE_MULTIPLIERS = {
        'Signature': (14, 3), 'Initial': (8, 2.5), 'Date': (10, 1.3),
        'Text': (8, 1.3), 'Checkbox': (2, 2), 'Name': (10, 1.3),
        'Title': (8, 1.3), 'Company': (8, 1.3), 'Email': (10, 1.3),
        'Date Input': (10, 1.3), 'Number': (8, 1.3), 'Currency': (8, 1.3),
        'Picklist': (10, 1.5)
    }

    TYPE_FLAG_MAP = {'d': 'Date Input', 'n': 'Number', 'm': 'Currency', 'p': 'Picklist'}

    def parse_size_overrides(suffix):
        w_match = re.search(r'W(\d+)', suffix, re.IGNORECASE)
        h_match = re.search(r'H(\d+)', suffix, re.IGNORECASE)
        return (
            int(w_match.group(1)) if w_match else None,
            int(h_match.group(1)) if h_match else None
        )

    def parse_flags(suffix):
        if not suffix:
            return False, False, False, 'Text'
        pre_fill = 'f' in suffix.lower()
        save_back = 's' in suffix.lower()
        required = 'r' in suffix.lower()
        field_type = 'Text'
        for flag, ftype in TYPE_FLAG_MAP.items():
            if flag in suffix.lower():
                field_type = ftype
                break
        return pre_fill, save_back, required, field_type

    def get_default_sizes(field_type, font_size):
        mw, mh = SIZE_MULTIPLIERS.get(field_type, (8, 1.3))
        return round(font_size * mw), round(font_size * mh)

    def parse_offsets(suffix):
        if not suffix:
            return 0, 0
        ox_match = re.search(r'OX(-?\d+)', suffix, re.IGNORECASE)
        oy_match = re.search(r'OY(-?\d+)', suffix, re.IGNORECASE)
        return (
            int(ox_match.group(1)) if ox_match else 0,
            int(oy_match.group(1)) if oy_match else 0
        )

    def baseline_to_top(y_pct, field_type, field_h, page_h):
        if page_h <= 0:
            return y_pct
        nudge = 2
        if field_type in ('Signature', 'Initial'):
            return y_pct - (field_h * 0.85 + nudge) / page_h * 100
        return y_pct - (field_h * 0.7 + nudge) / page_h * 100

    try:
        pdf_bytes = base64.b64decode(pdf_b64)
        doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to open PDF: {str(e)}")

    fields = []
    seen_tags = set()

    for page_idx in range(len(doc)):
        page = doc[page_idx]
        pw = page.rect.width
        ph = page.rect.height
        text_dict = page.get_text("dict")

        # Collect all spans, then concatenate text for cross-span matching
        span_list = []
        for block in text_dict.get("blocks", []):
            if block.get("type") != 0:
                continue
            for line in block.get("lines", []):
                for span in line.get("spans", []):
                    raw_text = span.get("text", "")
                    if not raw_text:
                        continue
                    span_list.append(span)

        # Build concatenated text with charMap (maps text index → span)
        full_text = ""
        char_map = []  # index → span reference
        for span in span_list:
            for ch in span.get("text", ""):
                char_map.append(span)
            full_text += span.get("text", "")

        # Whitespace normalization inside [[ ]] brackets
        clean_text = full_text
        clean_to_orig = list(range(len(full_text)))
        if "[" in full_text:
            depth = 0
            new_chars = []
            new_map = []
            for i, ch in enumerate(full_text):
                if ch == '[':
                    depth += 1; new_chars.append(ch); new_map.append(i)
                elif ch == ']':
                    depth -= 1; new_chars.append(ch); new_map.append(i)
                elif depth >= 2 and ch == ' ':
                    pass
                else:
                    new_chars.append(ch); new_map.append(i)
            clean_text = ''.join(new_chars)
            clean_to_orig = new_map

        def span_at(clean_idx):
            orig_idx = clean_to_orig[clean_idx] if clean_idx < len(clean_to_orig) else len(char_map) - 1
            return char_map[orig_idx] if orig_idx < len(char_map) else None

        def span_coords(sp):
            origin = sp.get("origin", (0, 0))
            fs = max(sp.get("size", 12), 12)
            xp = (origin[0] / pw) * 100 if pw > 0 else 0
            yp = (origin[1] / ph) * 100 if ph > 0 else 0
            return xp, yp, fs

        # Pass 1: WSM signing tags
        matched_orig_ranges = []
        for m in SIGNING_REGEX.finditer(clean_text):
            sp = span_at(m.start())
            if not sp:
                continue
            x_pct, y_pct, font_size = span_coords(sp)
            tag_key = m.group(1).upper()
            signer_num = int(m.group(2))
            suffix = m.group(3) or ""
            field_type = SIGNING_TYPE_MAP.get(tag_key, 'Text')

            dedup_key = f"{page_idx}_{tag_key}{signer_num}_{x_pct:.1f}_{y_pct:.1f}"
            if dedup_key in seen_tags:
                continue
            seen_tags.add(dedup_key)

            orig_start = clean_to_orig[m.start()]
            orig_end = clean_to_orig[m.end() - 1] + 1
            matched_orig_ranges.append((orig_start, orig_end))

            w_override, h_override = parse_size_overrides(suffix)
            def_w, def_h = get_default_sizes(field_type, font_size)
            final_h = h_override or def_h
            ox, oy = parse_offsets(suffix)
            required = 'r' in suffix.lower() if suffix else False

            fields.append({
                "pageNumber": page_idx + 1,
                "xPosition": round(x_pct + (ox / pw * 100 if pw > 0 else 0), 2),
                "yPosition": round(baseline_to_top(y_pct, field_type, final_h, ph) + (oy / ph * 100 if ph > 0 else 0), 2),
                "width": w_override or def_w,
                "height": final_h,
                "fieldType": field_type,
                "signerNumber": signer_num,
                "required": required,
                "label": f"{field_type} (Signer {signer_num})",
                "sourceField": None,
                "preFill": False,
                "saveBack": False
            })

        # Pass 2: WSM text tabs (skip ranges matched by signing tags)
        for m in TEXT_TAB_REGEX.finditer(clean_text):
            field_name = m.group(1)
            suffix = m.group(2) or ""
            if re.match(r'^(?:S|I|N|D|T|CO|EML)\d+$', field_name, re.IGNORECASE):
                continue
            orig_start = clean_to_orig[m.start()]
            orig_end = clean_to_orig[m.end() - 1] + 1
            if any(orig_start < re and orig_end > rs for rs, re in matched_orig_ranges):
                continue

            sp = span_at(m.start())
            if not sp:
                continue
            x_pct, y_pct, font_size = span_coords(sp)

            dedup_key = f"{page_idx}_{field_name}_{x_pct:.1f}_{y_pct:.1f}"
            if dedup_key in seen_tags:
                continue
            seen_tags.add(dedup_key)

            pre_fill, save_back, required, field_type = parse_flags(suffix)
            w_override, h_override = parse_size_overrides(suffix)
            def_w, def_h = get_default_sizes(field_type, font_size)
            tt_h = h_override or def_h
            tt_ox, tt_oy = parse_offsets(suffix)

            signer_num = 0
            trailing_match = re.match(r'.*_(\d)$', field_name)
            if trailing_match:
                signer_num = int(trailing_match.group(1))

            fields.append({
                "pageNumber": page_idx + 1,
                "xPosition": round(x_pct + (tt_ox / pw * 100 if pw > 0 else 0), 2),
                "yPosition": round(baseline_to_top(y_pct, field_type, tt_h, ph) + (tt_oy / ph * 100 if ph > 0 else 0), 2),
                "width": w_override or def_w,
                "height": tt_h,
                "fieldType": field_type,
                "signerNumber": signer_num,
                "required": required,
                "label": field_name,
                "sourceField": field_name,
                "preFill": pre_fill,
                "saveBack": save_back
            })

        # Pass 3: DocuSign signing anchors (\s1\, \i2\, etc.)
        if "\\" in full_text:
            for m in DS_SIGNING_REGEX.finditer(full_text):
                sp = char_map[m.start()] if m.start() < len(char_map) else None
                if not sp:
                    continue
                x_pct, y_pct, font_size = span_coords(sp)
                ds_key = m.group(1).lower()
                signer_num = int(m.group(2))
                field_type = DS_SIGNING_TYPE_MAP.get(ds_key, 'Text')

                dedup_key = f"{page_idx}_ds_{ds_key}{signer_num}_{x_pct:.1f}_{y_pct:.1f}"
                if dedup_key in seen_tags:
                    continue
                seen_tags.add(dedup_key)

                def_w, def_h = get_default_sizes(field_type, font_size)
                fields.append({
                    "pageNumber": page_idx + 1,
                    "xPosition": round(x_pct, 2),
                    "yPosition": round(baseline_to_top(y_pct, field_type, def_h, ph), 2),
                    "width": def_w,
                    "height": def_h,
                    "fieldType": field_type,
                    "signerNumber": signer_num,
                    "required": field_type in ('Signature', 'Initial'),
                    "label": field_type,
                    "sourceField": None,
                    "preFill": False,
                    "saveBack": False
                })

            for m in DS_TEXT_TAB_REGEX.finditer(full_text):
                sp = char_map[m.start()] if m.start() < len(char_map) else None
                if not sp:
                    continue
                x_pct, y_pct, font_size = span_coords(sp)
                ds_label = m.group(1)
                signer_num = int(m.group(2))
                ds_type_key = m.group(3).lower()
                field_type = DS_TEXT_TYPE_MAP.get(ds_type_key, 'Text')

                dedup_key = f"{page_idx}_ds_{ds_label}_{signer_num}_{x_pct:.1f}_{y_pct:.1f}"
                if dedup_key in seen_tags:
                    continue
                seen_tags.add(dedup_key)

                label = re.sub(r'([A-Z])', r' \1', ds_label).strip().title()
                def_w, def_h = get_default_sizes(field_type, font_size)
                fields.append({
                    "pageNumber": page_idx + 1,
                    "xPosition": round(x_pct, 2),
                    "yPosition": round(baseline_to_top(y_pct, field_type, def_h, ph), 2),
                    "width": def_w,
                    "height": def_h,
                    "fieldType": field_type,
                    "signerNumber": signer_num,
                    "required": False,
                    "label": label,
                    "sourceField": None,
                    "preFill": False,
                    "saveBack": False
                })

    total_pages = len(doc)
    doc.close()
    logger.info("extract-pdf-tags: fields=%d pages=%d duration=%.2fs", len(fields), total_pages, time.time() - start)
    return {"fields": fields, "totalPages": total_pages, "success": True}


WIDGET_TYPE_MAP = {
    0: "Unknown",
    1: "Button",
    2: "Checkbox",
    3: "Combobox",
    4: "Listbox",
    5: "RadioButton",
    6: "Signature",
    7: "Text",
}


@app.post("/detect-form-fields")
async def detect_form_fields(request: Request):
    _verify_api_key(request)
    body = await request.json()
    start = time.time()

    pdf_b64 = body.get("pdf_base64")
    if not pdf_b64:
        raise HTTPException(status_code=400, detail="Missing pdf_base64")
    if len(pdf_b64) > MAX_PDF_BASE64_BYTES:
        raise HTTPException(status_code=400, detail=f"PDF too large ({len(pdf_b64)} bytes, max {MAX_PDF_BASE64_BYTES})")

    try:
        pdf_bytes = base64.b64decode(pdf_b64)
        doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to open PDF: {str(e)}")

    fields = []
    for page_idx in range(len(doc)):
        page = doc[page_idx]
        pw = page.rect.width
        ph = page.rect.height

        for widget in page.widgets():
            if not widget.field_name:
                continue
            rect = widget.rect
            field_type = WIDGET_TYPE_MAP.get(widget.field_type, "Unknown")

            field = {
                "fieldName": widget.field_name,
                "fieldType": field_type,
                "pageNumber": page_idx + 1,
                "xPercent": round((rect.x0 / pw) * 100, 2) if pw > 0 else 0,
                "yPercent": round((rect.y0 / ph) * 100, 2) if ph > 0 else 0,
                "width": round(rect.width, 2),
                "height": round(rect.height, 2),
                "currentValue": widget.field_value or "",
            }

            if widget.field_type == 3 and widget.choice_values:
                field["options"] = [v[0] if isinstance(v, (list, tuple)) else v for v in widget.choice_values]

            fields.append(field)

            if len(fields) >= MAX_FORM_FIELDS:
                break
        if len(fields) >= MAX_FORM_FIELDS:
            break

    # Scan text for {{ merge_tags }}
    import re
    merge_tag_pattern = re.compile(r'\{\{\s*(.+?)\s*\}\}')
    merge_tags = set()
    for page_idx in range(len(doc)):
        page = doc[page_idx]
        text = page.get_text()
        for match in merge_tag_pattern.finditer(text):
            merge_tags.add(match.group(1).strip())

    doc.close()
    logger.info("detect-form-fields: fields=%d merge_tags=%d duration=%.2fs", len(fields), len(merge_tags), time.time() - start)
    return {
        "fields": fields,
        "hasFormFields": len(fields) > 0,
        "mergeTags": sorted(merge_tags),
        "hasMergeTags": len(merge_tags) > 0,
        "success": True,
    }


@app.post("/fill-pdf-form")
async def fill_pdf_form(request: Request):
    _verify_api_key(request)
    body = await request.json()
    start = time.time()

    pdf_b64 = body.get("pdf_base64")
    field_values = body.get("field_values", {})

    if not pdf_b64:
        raise HTTPException(status_code=400, detail="Missing pdf_base64")
    if len(pdf_b64) > MAX_PDF_BASE64_BYTES:
        raise HTTPException(status_code=400, detail=f"PDF too large ({len(pdf_b64)} bytes, max {MAX_PDF_BASE64_BYTES})")
    if not field_values:
        raise HTTPException(status_code=400, detail="Missing field_values")

    try:
        pdf_bytes = base64.b64decode(pdf_b64)
        doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to open PDF: {str(e)}")

    filled_count = 0
    for page in doc:
        for widget in page.widgets():
            if widget.field_name in field_values:
                val = field_values[widget.field_name]
                if widget.field_type == 2:
                    widget.field_value = bool(val)
                else:
                    widget.field_value = str(val)
                widget.update()
                filled_count += 1

    doc.bake()

    output = io.BytesIO()
    doc.save(output)
    doc.close()

    pdf_b64_out = base64.b64encode(output.getvalue()).decode("utf-8")
    logger.info("fill-pdf-form: filled=%d duration=%.2fs", filled_count, time.time() - start)
    return {"pdf_base64": pdf_b64_out, "filled_count": filled_count, "success": True}


@app.post("/render-pages")
async def render_pages(request: Request):
    _verify_api_key(request)
    body = await request.json()
    start = time.time()

    pdf_b64 = body.get("pdf_base64")
    dpi = min(body.get("dpi", 150), MAX_RENDER_DPI)
    requested_pages = body.get("pages")

    if not pdf_b64:
        raise HTTPException(status_code=400, detail="Missing pdf_base64")
    if len(pdf_b64) > MAX_PDF_BASE64_BYTES:
        raise HTTPException(status_code=400, detail=f"PDF too large ({len(pdf_b64)} bytes, max {MAX_PDF_BASE64_BYTES})")

    try:
        pdf_bytes = base64.b64decode(pdf_b64)
        doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to open PDF: {str(e)}")

    total_pages = len(doc)

    if requested_pages:
        page_nums = [p for p in requested_pages if 1 <= p <= total_pages]
    else:
        page_nums = list(range(1, min(total_pages + 1, MAX_RENDER_PAGES + 1)))

    scale = dpi / 72.0
    matrix = fitz.Matrix(scale, scale)
    pages_out = []

    for pn in page_nums:
        page = doc[pn - 1]
        pixmap = page.get_pixmap(matrix=matrix)
        img_bytes = pixmap.tobytes("png")
        pages_out.append({
            "pageNumber": pn,
            "imageBase64": base64.b64encode(img_bytes).decode("utf-8"),
            "widthPts": round(page.rect.width, 2),
            "heightPts": round(page.rect.height, 2),
        })

    doc.close()
    logger.info("render-pages: pages=%d dpi=%d duration=%.2fs", len(pages_out), dpi, time.time() - start)
    return {"pages": pages_out, "totalPages": total_pages, "success": True}


# ── SF OAuth helper ────────────────────────────────────────────────────

def _sf_authenticate(instance_url: str, client_id: str, client_secret: str):
    resp = http_requests.post(
        f"{instance_url}/services/oauth2/token",
        data={
            "grant_type": "client_credentials",
            "client_id": client_id,
            "client_secret": client_secret,
        },
        timeout=30,
    )
    if resp.status_code != 200:
        raise HTTPException(
            status_code=401,
            detail=f"SF OAuth failed (HTTP {resp.status_code}): {resp.text[:300]}",
        )
    return resp.json()["access_token"]


def _sf_helpers(instance_url, access_token, api_version="63.0"):
    base_url = f"{instance_url}/services/data/v{api_version}"
    auth_headers = {"Authorization": f"Bearer {access_token}"}

    def query(soql):
        resp = http_requests.get(
            f"{base_url}/query", headers=auth_headers,
            params={"q": soql}, timeout=30,
        )
        if resp.status_code != 200:
            raise Exception(f"SF query failed (HTTP {resp.status_code}): {resp.text[:500]}")
        return resp.json().get("records", [])

    def fetch_file(content_document_id):
        records = query(
            f"SELECT Id FROM ContentVersion "
            f"WHERE ContentDocumentId = '{content_document_id}' AND IsLatest = true LIMIT 1"
        )
        if not records:
            raise Exception(f"No ContentVersion for ContentDocumentId {content_document_id}")
        cv_id = records[0]["Id"]
        resp = http_requests.get(
            f"{base_url}/sobjects/ContentVersion/{cv_id}/VersionData",
            headers=auth_headers, timeout=60,
        )
        if resp.status_code != 200:
            raise Exception(f"File download failed (HTTP {resp.status_code})")
        return resp.content

    def create_content_version(file_bytes, filename, parent_id, mime_type="application/pdf"):
        import json as _json
        import os as _os
        entity = {
            "Title": _os.path.splitext(filename)[0],
            "PathOnClient": filename,
            "FirstPublishLocationId": parent_id,
        }
        resp = http_requests.post(
            f"{base_url}/sobjects/ContentVersion",
            headers=auth_headers,
            files={
                "entity_content": (None, _json.dumps(entity), "application/json"),
                "VersionData": (filename, file_bytes, mime_type),
            },
            timeout=120,
        )
        if resp.status_code not in (200, 201):
            raise Exception(f"ContentVersion create failed (HTTP {resp.status_code}): {resp.text[:500]}")
        cv_id = resp.json()["id"]
        cv_records = query(
            f"SELECT ContentDocumentId FROM ContentVersion WHERE Id = '{cv_id}' LIMIT 1"
        )
        content_doc_id = cv_records[0]["ContentDocumentId"] if cv_records else None
        return cv_id, content_doc_id

    return query, fetch_file, create_content_version


# ── Word XML helpers (ported from modal_api.py) ───────────────────────

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
XML_NS = "http://www.w3.org/XML/1998/namespace"

DOCUSIGN_ANY_TAG = r'(?:&lt;|<)#\s*(?:&lt;|<)(?:Content|Conditional|EndConditional|Signature|TableRow|EndTableRow|/TableRow)\b[^#]*#(?:&gt;|>)'
CONTENT_PATTERN = r'(?:&lt;|<)#\s*(?:&lt;|<)Content\s+Select="([^"]+)"\s*/(?:&gt;|>)\s*#(?:&gt;|>)'
CONDITIONAL_PATTERN = r'(?:&lt;|<)#\s*(?:&lt;|<)Conditional\s+Select="([^"]+)"\s+Match="([^"]+)"\s*/(?:&gt;|>)\s*#(?:&gt;|>)'
END_CONDITIONAL_PATTERN = r'(?:&lt;|<)#\s*(?:&lt;|<)EndConditional\s*/(?:&gt;|>)\s*#(?:&gt;|>)'
SIGNATURE_PATTERN = r'(?:&lt;|<)#\s*(?:&lt;|<)Signature\s+Placeholder="\\([^"\\]+)\\"\s+Hidden="[^"]*"\s*/(?:&gt;|>)\s*#(?:&gt;|>)'
TABLEROW_PATTERN = r'(?:&lt;|<)#\s*(?:&lt;|<)TableRow\s+Select="([^"]+)"\s*/?(?:&gt;|>)\s*#(?:&gt;|>)'
END_TABLEROW_PATTERN = (
    r'(?:(?:&lt;|<)#\s*(?:&lt;|<)/TableRow\s*(?:&gt;|>)\s*#(?:&gt;|>))'
    r'|(?:(?:&lt;|<)#\s*(?:&lt;|<)EndTableRow\s*/(?:&gt;|>)\s*#(?:&gt;|>))'
)
TEXT_TAB_PATTERN = r'\\([a-zA-Z][a-zA-Z0-9]*)_(\d+)_text\\'
HIDDEN_FIELD_PATTERN = re.compile(r'\[\[[a-zA-Z0-9_.\-]+\]\]')


def _et_tostring_preserve_decl(root, original_xml):
    from lxml import etree as ET
    result = ET.tostring(root, encoding="unicode")
    if isinstance(original_xml, str) and original_xml.lstrip().startswith("<?xml"):
        decl_end = original_xml.find("?>") + 2
        decl = original_xml[:decl_end]
        nl = "\n" if original_xml[decl_end : decl_end + 1] in ("\n", "\r") else ""
        return decl + nl + result
    return result


def _merge_bracket_runs(root):
    from lxml import etree as ET
    w_p_tag = f"{{{W_NS}}}p"
    w_r_tag = f"{{{W_NS}}}r"
    w_t_tag = f"{{{W_NS}}}t"
    bracket_pattern = re.compile(r"\[\[.+?\]\]")
    for para in root.iter(w_p_tag):
        run_texts = []
        for child in para:
            if child.tag == w_r_tag:
                for t in child.findall(w_t_tag):
                    run_texts.append({"run": child, "t": t, "text": t.text or ""})
        if len(run_texts) < 2:
            continue
        full_text = "".join(rt["text"] for rt in run_texts)
        if not bracket_pattern.search(full_text):
            continue
        positions = []
        offset = 0
        for i, rt in enumerate(run_texts):
            length = len(rt["text"])
            positions.append({"start": offset, "end": offset + length, "index": i})
            offset += length
        for m in bracket_pattern.finditer(full_text):
            tag_start, tag_end = m.start(), m.end()
            affected = [p["index"] for p in positions if p["start"] < tag_end and p["end"] > tag_start]
            if len(affected) > 1:
                first = affected[0]
                merged = "".join(run_texts[i]["text"] for i in affected)
                run_texts[first]["t"].text = merged
                run_texts[first]["t"].set(f"{{{W_NS}}}space", "preserve")
                for i in affected[1:]:
                    run_texts[i]["t"].text = ""


def hide_placeholder_fields(xml_content, anchor_color="FFFFFF"):
    from lxml import etree as ET
    import copy
    try:
        root = ET.fromstring(xml_content.encode("utf-8") if isinstance(xml_content, str) else xml_content)
    except ET.XMLSyntaxError:
        return xml_content
    _merge_bracket_runs(root)
    w_r_tag = f"{{{W_NS}}}r"
    w_t_tag = f"{{{W_NS}}}t"
    w_rPr_tag = f"{{{W_NS}}}rPr"
    w_color_tag = f"{{{W_NS}}}color"
    w_val_attr = f"{{{W_NS}}}val"
    w_sz_tag = f"{{{W_NS}}}sz"
    w_szCs_tag = f"{{{W_NS}}}szCs"
    runs_to_split = []
    for t_elem in root.iter(w_t_tag):
        if t_elem.text and HIDDEN_FIELD_PATTERN.search(t_elem.text):
            run = t_elem.getparent()
            if run is None or run.tag != w_r_tag:
                continue
            if HIDDEN_FIELD_PATTERN.sub("", t_elem.text).strip():
                runs_to_split.append((run, t_elem))
            else:
                rPr = run.find(w_rPr_tag)
                if rPr is None:
                    rPr = ET.Element(w_rPr_tag)
                    run.insert(0, rPr)
                color_elem = rPr.find(w_color_tag)
                if color_elem is None:
                    color_elem = ET.SubElement(rPr, w_color_tag)
                color_elem.set(w_val_attr, anchor_color)
                for tag in (w_sz_tag, w_szCs_tag):
                    el = rPr.find(tag)
                    if el is None:
                        el = ET.SubElement(rPr, tag)
                    el.set(w_val_attr, "2")
    for run, t_elem in runs_to_split:
        para = run.getparent()
        if para is None:
            continue
        run_index = list(para).index(run)
        text = t_elem.text
        parts = HIDDEN_FIELD_PATTERN.split(text)
        placeholders = HIDDEN_FIELD_PATTERN.findall(text)
        orig_rPr = run.find(w_rPr_tag)
        new_runs = []
        for i, part in enumerate(parts):
            if part:
                r = ET.Element(w_r_tag)
                if orig_rPr is not None:
                    r.append(copy.deepcopy(orig_rPr))
                t = ET.SubElement(r, w_t_tag)
                t.text = part
                t.set(f"{{{XML_NS}}}space", "preserve")
                new_runs.append(r)
            if i < len(placeholders):
                r = ET.Element(w_r_tag)
                rPr = copy.deepcopy(orig_rPr) if orig_rPr is not None else ET.Element(w_rPr_tag)
                color_elem = rPr.find(w_color_tag)
                if color_elem is None:
                    color_elem = ET.SubElement(rPr, w_color_tag)
                color_elem.set(w_val_attr, anchor_color)
                for tag in (w_sz_tag, w_szCs_tag):
                    el = rPr.find(tag)
                    if el is None:
                        el = ET.SubElement(rPr, tag)
                    el.set(w_val_attr, "2")
                r.insert(0, rPr)
                t = ET.SubElement(r, w_t_tag)
                t.text = placeholders[i]
                t.set(f"{{{XML_NS}}}space", "preserve")
                new_runs.append(r)
        para.remove(run)
        for j, new_run in enumerate(new_runs):
            para.insert(run_index + j, new_run)
    return _et_tostring_preserve_decl(root, xml_content)


def normalize_xml_runs(xml_content):
    from lxml import etree as ET
    try:
        root = ET.fromstring(xml_content.encode("utf-8") if isinstance(xml_content, str) else xml_content)
    except ET.XMLSyntaxError:
        return xml_content
    w_p = f"{{{W_NS}}}p"
    w_r = f"{{{W_NS}}}r"
    w_t = f"{{{W_NS}}}t"
    for para in root.iter(w_p):
        run_texts = []
        for run in para.iter(w_r):
            for t_elem in run.iter(w_t):
                run_texts.append((run, t_elem))
        if not run_texts:
            continue
        full_text = "".join((t.text or "") for _, t in run_texts)
        if not re.search(DOCUSIGN_ANY_TAG, full_text):
            continue
        run_positions = []
        offset = 0
        for i, (_, t) in enumerate(run_texts):
            text = t.text or ""
            run_positions.append((offset, offset + len(text), i))
            offset += len(text)
        merge_groups = []
        for m in re.finditer(DOCUSIGN_ANY_TAG, full_text):
            tag_start, tag_end = m.start(), m.end()
            affected = [i for (s, e, i) in run_positions if s < tag_end and e > tag_start]
            if len(affected) > 1:
                merge_groups.append((affected[0], affected[-1]))
        if not merge_groups:
            continue
        for first_idx, last_idx in reversed(merge_groups):
            merged_text = "".join((run_texts[i][1].text or "") for i in range(first_idx, last_idx + 1))
            run_texts[first_idx][1].text = merged_text
            run_texts[first_idx][1].set(f"{{{XML_NS}}}space", "preserve")
            for i in range(first_idx + 1, last_idx + 1):
                run_texts[i][1].text = ""
    return _et_tostring_preserve_decl(root, xml_content)


def process_tablerow_sections(xml_content):
    section_pattern = (
        r"(" + TABLEROW_PATTERN + r")(.*?)(" + END_TABLEROW_PATTERN + r")"
    )
    def replace_section(match):
        xpath = match.group(2)
        parts = xpath.split("//")
        collection = parts[-1].strip("/").split("/")[-1] if len(parts) > 1 else xpath.strip("/").split("/")[-1]
        inner = match.group(3)
        def replace_inner_content(m):
            field_path = m.group(1).lstrip("./").strip("/")
            return "{{ item." + ".".join(field_path.split("/")) + " }}"
        inner = re.sub(CONTENT_PATTERN, replace_inner_content, inner)
        return "{%tr for item in " + collection + " %}" + inner + "{%tr endfor %}"
    return re.sub(section_pattern, replace_section, xml_content, flags=re.DOTALL)


# ── /generate-documents — unified doc gen endpoint ─────────────────────

INSTANCE_URL_RE = re.compile(r"^https://[a-zA-Z0-9.-]+\.my\.salesforce\.com$")
SF_ID_RE = re.compile(r"^[a-zA-Z0-9]{15,18}$")


@app.post("/generate-documents")
async def generate_documents(request: Request):
    _verify_api_key(request)
    body = await request.json()
    start = time.time()

    import json as _json
    import subprocess
    import tempfile
    import zipfile
    import os
    import copy
    from pathlib import Path
    from lxml import etree as ET

    # Optional per-request Microsoft Graph config for DOCX->PDF conversion.
    # Salesforce merges these into the request body from the External Credential
    # principal (same mechanism as sf_client_id/sf_client_secret) — never as
    # custom headers. Each key falls back to its env var, so env-configured
    # deployments stay unchanged when the body keys are absent.
    ms_config = {
        "tenant": body.get("ms_tenant_id") or os.environ.get("MS_TENANT_ID"),
        "client": body.get("ms_client_id") or os.environ.get("MS_CLIENT_ID"),
        "secret": body.get("ms_client_secret") or os.environ.get("MS_CLIENT_SECRET"),
        "user": body.get("ms_user_id") or os.environ.get("MS_USER_ID"),
    }

    template_files_meta = body.get("template_files", [])
    merge_data = body.get("merge_data", {})
    source_record_id = body.get("source_record_id")
    instance_url = body.get("instance_url", "")
    api_version = body.get("api_version", "63.0")
    merge_into_one = body.get("merge_into_one", False)
    output_filename = body.get("output_filename", "Merged Document.pdf")
    sf_client_id = body.get("sf_client_id", "")
    sf_client_secret = body.get("sf_client_secret", "")

    if not template_files_meta:
        raise HTTPException(status_code=400, detail="Missing template_files")
    if not source_record_id:
        raise HTTPException(status_code=400, detail="Missing source_record_id")
    if not INSTANCE_URL_RE.match(instance_url):
        raise HTTPException(status_code=400, detail="Invalid or missing instance_url")
    if not sf_client_id or not sf_client_secret:
        raise HTTPException(status_code=401, detail="Missing SF credentials")
    if not SF_ID_RE.match(source_record_id):
        raise HTTPException(status_code=400, detail=f"Invalid Salesforce ID: {source_record_id}")

    access_token = _sf_authenticate(instance_url, sf_client_id, sf_client_secret)
    sf_query, sf_fetch_file, sf_create_cv = _sf_helpers(instance_url, access_token, api_version)

    def convert_to_pdf(input_path, tmpdir):
        ext = str(input_path).rsplit(".", 1)[-1].lower()
        if ext in ("docx", "doc") and ms_config.get("client"):
            try:
                return _convert_via_graph(input_path)
            except Exception as graph_err:
                logger.warning("Graph API conversion failed, falling back to LibreOffice: %s", graph_err)
        result = subprocess.run(
            ["soffice", "--headless", "--convert-to", "pdf", "--outdir", tmpdir, str(input_path)],
            check=True, timeout=90, capture_output=True, text=True,
        )
        pdf_path = Path(tmpdir) / (input_path.stem + ".pdf")
        if not pdf_path.exists():
            raise Exception(f"LibreOffice PDF conversion failed: {result.stderr}")
        with open(pdf_path, "rb") as f:
            return f.read()

    def _convert_via_graph(input_path):
        ms_tenant = ms_config.get("tenant")
        ms_client = ms_config.get("client")
        ms_secret = ms_config.get("secret")
        ms_user = ms_config.get("user")
        if not all([ms_tenant, ms_client, ms_secret, ms_user]):
            raise Exception("Incomplete Microsoft Graph config (need tenant, client, secret, user)")

        token_resp = http_requests.post(
            f"https://login.microsoftonline.com/{ms_tenant}/oauth2/v2.0/token",
            data={
                "grant_type": "client_credentials",
                "client_id": ms_client,
                "client_secret": ms_secret,
                "scope": "https://graph.microsoft.com/.default",
            },
            timeout=30,
        )
        if token_resp.status_code != 200:
            raise Exception(f"Graph auth failed: {token_resp.text[:200]}")
        ms_token = token_resp.json()["access_token"]
        ms_headers = {"Authorization": f"Bearer {ms_token}"}

        with open(input_path, "rb") as f:
            docx_bytes = f.read()

        filename = f"docgen-{time.time():.0f}.docx"
        upload_resp = http_requests.put(
            f"https://graph.microsoft.com/v1.0/users/{ms_user}/drive/root:/docgen-temp/{filename}:/content",
            headers={**ms_headers, "Content-Type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document"},
            data=docx_bytes,
            timeout=60,
        )
        if upload_resp.status_code not in (200, 201):
            raise Exception(f"Graph upload failed: {upload_resp.text[:200]}")

        item = upload_resp.json()
        item_id = item["id"]
        drive_id = item["parentReference"]["driveId"]

        import time as _time
        _time.sleep(1)

        pdf_resp = http_requests.get(
            f"https://graph.microsoft.com/v1.0/drives/{drive_id}/items/{item_id}/content?format=pdf",
            headers=ms_headers,
            timeout=120,
        )

        http_requests.delete(
            f"https://graph.microsoft.com/v1.0/drives/{drive_id}/items/{item_id}",
            headers=ms_headers,
            timeout=15,
        )

        if pdf_resp.status_code != 200 or len(pdf_resp.content) < 100:
            raise Exception(f"Graph PDF conversion failed: HTTP {pdf_resp.status_code}")

        logger.info("Graph API conversion: %s → %d bytes PDF", input_path.name if hasattr(input_path, 'name') else input_path, len(pdf_resp.content))
        return pdf_resp.content

    def _convert_docusign_to_jinja2(docx_path):
        with zipfile.ZipFile(docx_path, "r") as zf:
            has_docusign = False
            for name in zf.namelist():
                if name.startswith("word/") and name.endswith(".xml"):
                    content = zf.read(name).decode("utf-8", errors="ignore")
                    if re.search(DOCUSIGN_ANY_TAG, content) or re.search(TEXT_TAB_PATTERN, content):
                        has_docusign = True
                        break
            if not has_docusign:
                return docx_path

        temp_dir = docx_path.parent / "docx_temp"
        temp_dir.mkdir(exist_ok=True)
        with zipfile.ZipFile(docx_path, "r") as zip_ref:
            zip_ref.extractall(temp_dir)

        def replace_field(match):
            xpath = match.group(1)
            parts = xpath.strip("/").split("/")
            return "{{ " + (".".join(parts[1:]) if len(parts) > 1 else parts[0]) + " }}"

        def replace_conditional(match):
            xpath = match.group(1)
            parts = xpath.strip("/").split("/")
            key = ".".join(parts[1:]) if len(parts) > 1 else parts[0]
            return '{{% if {k} == "{v}" %}}'.format(k=key, v=match.group(2))

        def replace_end_conditional(match):
            return "{% endif %}"

        def replace_signature(match):
            return "[[" + match.group(1).upper() + "]]"

        def replace_text_tab(match):
            return f"[[{match.group(1)}_{match.group(2)}]]"

        def _split_colon_value_runs(xml_content):
            colon_value_re = re.compile(r"^(.*\S)\s*(:)\s*(\{\{.*\}\}.*)$", re.DOTALL)
            try:
                root = ET.fromstring(xml_content.encode("utf-8") if isinstance(xml_content, str) else xml_content)
            except ET.XMLSyntaxError:
                return xml_content
            w_p = f"{{{W_NS}}}p"
            w_r = f"{{{W_NS}}}r"
            w_t = f"{{{W_NS}}}t"
            w_rPr = f"{{{W_NS}}}rPr"
            w_tab = f"{{{W_NS}}}tab"
            w_tabs = f"{{{W_NS}}}tabs"
            w_pPr = f"{{{W_NS}}}pPr"
            for para in root.iter(w_p):
                pPr = para.find(w_pPr)
                if pPr is None or pPr.find(w_tabs) is None:
                    continue
                runs = list(para)
                for run in runs:
                    if run.tag != w_r:
                        continue
                    t_elem = run.find(w_t)
                    if t_elem is None or not t_elem.text:
                        continue
                    m = colon_value_re.match(t_elem.text)
                    if not m:
                        continue
                    label, colon, value = m.group(1), m.group(2), m.group(3)
                    orig_rPr = run.find(w_rPr)
                    run_idx = list(para).index(run)
                    t_elem.text = label
                    t_elem.set(f"{{{XML_NS}}}space", "preserve")
                    tab_run = ET.Element(w_r)
                    if orig_rPr is not None:
                        tab_run.append(copy.deepcopy(orig_rPr))
                    ET.SubElement(tab_run, w_tab)
                    tab_t = ET.SubElement(tab_run, w_t)
                    tab_t.text = f"{colon} {value}"
                    tab_t.set(f"{{{XML_NS}}}space", "preserve")
                    para.insert(run_idx + 1, tab_run)
                    break
            return _et_tostring_preserve_decl(root, xml_content)

        def process_xml_file(xml_path):
            if not xml_path.exists():
                return
            with open(xml_path, "r", encoding="utf-8") as f:
                content = f.read()
            if not re.search(DOCUSIGN_ANY_TAG, content) and not re.search(TEXT_TAB_PATTERN, content):
                return
            content = normalize_xml_runs(content)
            modified = process_tablerow_sections(content)
            modified = re.sub(CONTENT_PATTERN, replace_field, modified)
            modified = re.sub(CONDITIONAL_PATTERN, replace_conditional, modified)
            modified = re.sub(END_CONDITIONAL_PATTERN, replace_end_conditional, modified)
            modified = re.sub(SIGNATURE_PATTERN, replace_signature, modified)
            modified = re.sub(TEXT_TAB_PATTERN, replace_text_tab, modified)
            modified = _split_colon_value_runs(modified)
            with open(xml_path, "w", encoding="utf-8") as f:
                f.write(modified)

        process_xml_file(temp_dir / "word" / "document.xml")
        word_dir = temp_dir / "word"
        if word_dir.exists():
            for hf in list(word_dir.glob("header*.xml")) + list(word_dir.glob("footer*.xml")):
                process_xml_file(hf)

        output_path = docx_path.parent / "converted_template.docx"
        with zipfile.ZipFile(output_path, "w", zipfile.ZIP_DEFLATED) as zip_out:
            for walk_root, dirs, files in os.walk(temp_dir):
                for file in files:
                    file_path = Path(walk_root) / file
                    zip_out.write(file_path, file_path.relative_to(temp_dir))
        return output_path

    def process_word_template(docx_bytes, merge_data_inner, anchor_tag_color, tmpdir):
        from docxtpl import DocxTemplate
        import jinja2
        work_dir = Path(tmpdir) / "word_work"
        work_dir.mkdir(exist_ok=True)
        input_path = work_dir / "template.docx"
        with open(input_path, "wb") as f:
            f.write(docx_bytes)
        converted_path = _convert_docusign_to_jinja2(input_path)
        doc = DocxTemplate(converted_path)
        jinja_env = jinja2.Environment(autoescape=True, undefined=jinja2.ChainableUndefined)
        doc.render(merge_data_inner, jinja_env=jinja_env)
        rendered_path = work_dir / "rendered.docx"
        doc.save(rendered_path)
        hidden_dir = work_dir / "hidden_temp"
        hidden_dir.mkdir(exist_ok=True)
        with zipfile.ZipFile(rendered_path, "r") as zip_ref:
            zip_ref.extractall(hidden_dir)
        word_dir = hidden_dir / "word"
        xml_targets = [word_dir / "document.xml"]
        if word_dir.exists():
            xml_targets += list(word_dir.glob("header*.xml")) + list(word_dir.glob("footer*.xml"))
        for xml_file in xml_targets:
            if xml_file.exists():
                with open(xml_file, "r", encoding="utf-8") as f:
                    content = f.read()
                content = hide_placeholder_fields(content, anchor_tag_color)
                with open(xml_file, "w", encoding="utf-8") as f:
                    f.write(content)
        final_path = work_dir / "final.docx"
        with zipfile.ZipFile(final_path, "w", zipfile.ZIP_DEFLATED) as zip_out:
            for walk_root, dirs, files in os.walk(hidden_dir):
                for file in files:
                    file_path = Path(walk_root) / file
                    zip_out.write(file_path, file_path.relative_to(hidden_dir))
        return convert_to_pdf(final_path, str(work_dir))

    def process_rich_text_template(html_content, merge_data_inner, tmpdir):
        import jinja2
        env = jinja2.Environment(autoescape=True)
        template = env.from_string(html_content)
        rendered = template.render(**merge_data_inner)
        full_html = (
            '<!DOCTYPE html><html><head><meta charset="utf-8"><style>'
            "body { font-family: Arial, Helvetica, sans-serif; font-size: 12pt; margin: 1in; }"
            "table { border-collapse: collapse; width: 100%; }"
            "td, th { border: 1px solid #ccc; padding: 6px 8px; }"
            "img { max-width: 100%; }"
            "h1 { font-size: 20pt; } h2 { font-size: 16pt; } h3 { font-size: 14pt; }"
            "</style></head><body>" + rendered + "</body></html>"
        )
        html_path = Path(tmpdir) / "richtext.html"
        with open(html_path, "w", encoding="utf-8") as f:
            f.write(full_html)
        return convert_to_pdf(html_path, tmpdir)

    def process_pdf_template(pdf_bytes, tags_json, merge_data_inner, tmpdir):
        if not tags_json:
            return pdf_bytes
        mapping = _json.loads(tags_json) if isinstance(tags_json, str) else tags_json
        field_values = {}
        for form_field_name, sf_field_ref in mapping.items():
            if isinstance(sf_field_ref, str):
                sf_field = sf_field_ref.strip()
                if sf_field.startswith("{{") and sf_field.endswith("}}"):
                    sf_field = sf_field[2:-2].strip()
                if sf_field and sf_field in merge_data_inner:
                    field_values[form_field_name] = merge_data_inner[sf_field]
        if not field_values:
            return pdf_bytes
        doc = fitz.open(stream=pdf_bytes, filetype="pdf")
        for page in doc:
            for widget in page.widgets():
                if widget.field_name in field_values:
                    val = field_values[widget.field_name]
                    widget.field_value = bool(val) if widget.field_type == 2 else str(val)
                    widget.update()
        doc.bake()
        output = io.BytesIO()
        doc.save(output)
        doc.close()
        return output.getvalue()

    def process_excel_template(xlsx_bytes, merge_data_inner, file_tmpdir):
        # Render an .xlsx template against merge_data with xltpl (Jinja2 row/column
        # expansion + format preservation). ChainableUndefined matches the Word path:
        # missing merge fields render blank instead of raising. Returns native xlsx bytes.
        import jinja2
        import openpyxl
        from xltpl.writerx import BookWriter

        work_dir = Path(file_tmpdir) / "excel_work"
        work_dir.mkdir(exist_ok=True)
        input_path = work_dir / "template.xlsx"
        with open(input_path, "wb") as f:
            f.write(xlsx_bytes)

        # Source sheet names must be read before BookWriter loads (it moves the
        # template sheets into its own resource map; writer.workbook is the empty
        # output book). Render each source sheet to a same-named output sheet.
        sheet_names = openpyxl.load_workbook(input_path, read_only=True).sheetnames

        writer = BookWriter(str(input_path))
        writer.jinja_env.undefined = jinja2.ChainableUndefined  # missing vars -> blank (Word parity)
        writer.jinja_env.globals.update(merge_data_inner)

        # The payload doubles as routing (tpl_name/sheet_name) and the render context.
        payloads = [dict(merge_data_inner, tpl_name=n, sheet_name=n) for n in sheet_names]
        writer.render_book(payloads=payloads)

        rendered_path = work_dir / "rendered.xlsx"
        writer.save(str(rendered_path))
        with open(rendered_path, "rb") as f:
            return f.read()

    # Main processing
    template_files = []
    for fm in template_files_meta:
        template_files.append({
            "Id": fm.get("id", ""),
            "Name": fm.get("name", "document"),
            "File_Type__c": fm.get("file_type", ""),
            "Content_Document_Id__c": fm.get("content_document_id"),
            "Template_Tags__c": fm.get("template_tags"),
            "Anchor_Tag_Color__c": fm.get("anchor_tag_color"),
            "Html_Content__c": fm.get("html_content"),
        })

    generated_files = []
    with tempfile.TemporaryDirectory() as tmpdir:
        for idx, tf in enumerate(template_files):
            template_file_id = tf["Id"]
            name = tf.get("Name") or "document"
            file_type = (tf.get("File_Type__c") or "").strip()
            try:
                if file_type == "Rich Text":
                    html_content = tf.get("Html_Content__c")
                    if not html_content:
                        generated_files.append({"template_file_id": template_file_id, "error": "No HTML content"})
                        continue
                    file_tmpdir = os.path.join(tmpdir, f"rt_{idx}")
                    os.makedirs(file_tmpdir, exist_ok=True)
                    pdf_bytes = process_rich_text_template(html_content, merge_data, file_tmpdir)
                elif file_type == "PDF":
                    content_doc_id = tf.get("Content_Document_Id__c")
                    if not content_doc_id:
                        generated_files.append({"template_file_id": template_file_id, "error": "No PDF file"})
                        continue
                    pdf_bytes = sf_fetch_file(content_doc_id)
                    tags_json = tf.get("Template_Tags__c")
                    if tags_json:
                        file_tmpdir = os.path.join(tmpdir, f"pdf_{idx}")
                        os.makedirs(file_tmpdir, exist_ok=True)
                        pdf_bytes = process_pdf_template(pdf_bytes, tags_json, merge_data, file_tmpdir)
                elif file_type == "Excel":
                    # Native .xlsx output (generation only — not a signing format,
                    # and never folded into a merge_into_one combined PDF).
                    content_doc_id = tf.get("Content_Document_Id__c")
                    if not content_doc_id:
                        generated_files.append({"template_file_id": template_file_id, "error": "No .xlsx file"})
                        continue
                    xlsx_template_bytes = sf_fetch_file(content_doc_id)
                    file_tmpdir = os.path.join(tmpdir, f"excel_{idx}")
                    os.makedirs(file_tmpdir, exist_ok=True)
                    xlsx_bytes = process_excel_template(xlsx_template_bytes, merge_data, file_tmpdir)
                    cv_id, content_doc_id = sf_create_cv(
                        xlsx_bytes, f"{name}.xlsx", source_record_id,
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                    )
                    generated_files.append({
                        "template_file_id": template_file_id,
                        "content_version_id": cv_id,
                        "content_document_id": content_doc_id,
                    })
                    logger.info("Processed %s (Excel) → CV %s", name, cv_id)
                    continue
                else:
                    content_doc_id = tf.get("Content_Document_Id__c")
                    if not content_doc_id:
                        generated_files.append({"template_file_id": template_file_id, "error": "No .docx file"})
                        continue
                    docx_bytes = sf_fetch_file(content_doc_id)
                    anchor_color = tf.get("Anchor_Tag_Color__c") or "FFFFFF"
                    file_tmpdir = os.path.join(tmpdir, f"word_{idx}")
                    os.makedirs(file_tmpdir, exist_ok=True)
                    pdf_bytes = process_word_template(docx_bytes, merge_data, anchor_color, file_tmpdir)

                pdf_filename = f"{name}.pdf"
                cv_id, content_doc_id = sf_create_cv(pdf_bytes, pdf_filename, source_record_id)
                entry = {
                    "template_file_id": template_file_id,
                    "content_version_id": cv_id,
                    "content_document_id": content_doc_id,
                }
                if merge_into_one:
                    entry["_pdf_bytes"] = pdf_bytes
                generated_files.append(entry)
                logger.info("Processed %s (%s) → CV %s", name, file_type or "Word", cv_id)
            except Exception as e:
                logger.error("Error processing template file %s: %s", template_file_id, str(e))
                generated_files.append({"template_file_id": template_file_id, "error": str(e)})

    result = {"generated_files": generated_files}

    successful_pdfs = [gf for gf in generated_files if "error" not in gf and gf.get("_pdf_bytes")]
    if merge_into_one and len(successful_pdfs) > 1:
        try:
            merged_doc = fitz.open()
            for gf in successful_pdfs:
                doc = fitz.open(stream=gf["_pdf_bytes"], filetype="pdf")
                merged_doc.insert_pdf(doc)
                doc.close()
            merged_bytes = merged_doc.tobytes()
            merged_doc.close()
            cv_id, content_doc_id = sf_create_cv(merged_bytes, output_filename, source_record_id)
            result["merged_file"] = {
                "content_version_id": cv_id,
                "content_document_id": content_doc_id,
                "total_pages": fitz.open(stream=merged_bytes, filetype="pdf").page_count,
            }
            logger.info("Merged %d PDFs → CV %s", len(successful_pdfs), cv_id)
        except Exception as e:
            logger.error("PDF merge failed: %s", str(e))
            result["merged_file"] = {"error": str(e)}

    for gf in generated_files:
        gf.pop("_pdf_bytes", None)

    logger.info("generate-documents: files=%d duration=%.2fs", len(template_files), time.time() - start)
    return result


# ── /merge — SF-authenticated PDF merge ────────────────────────────────

@app.post("/merge")
async def merge_pdfs_authenticated(request: Request):
    _verify_api_key(request)
    body = await request.json()
    start = time.time()

    content_document_ids = body.get("content_document_ids", [])
    instance_url = body.get("instance_url", "")
    sf_client_id = body.get("sf_client_id", "")
    sf_client_secret = body.get("sf_client_secret", "")
    parent_record_id = body.get("parent_record_id")
    output_filename = body.get("output_filename", "Merged Document.pdf")
    api_version = body.get("api_version", "63.0")

    if not content_document_ids:
        raise HTTPException(status_code=400, detail="Missing content_document_ids")
    if not instance_url or not sf_client_id or not sf_client_secret:
        raise HTTPException(status_code=401, detail="Missing SF credentials")

    access_token = _sf_authenticate(instance_url, sf_client_id, sf_client_secret)
    sf_query, sf_fetch_file, sf_create_cv = _sf_helpers(instance_url, access_token, api_version)

    merged = fitz.open()
    page_counts = []

    for doc_id in content_document_ids:
        pdf_bytes = sf_fetch_file(doc_id)
        doc = fitz.open(stream=pdf_bytes, filetype="pdf")
        page_counts.append(len(doc))
        merged.insert_pdf(doc)
        doc.close()

    merged_bytes = merged.tobytes()
    merged.close()

    cv_id, content_doc_id = sf_create_cv(merged_bytes, output_filename, parent_record_id or "")
    total = sum(page_counts)
    logger.info("merge: docs=%d pages=%d duration=%.2fs", len(content_document_ids), total, time.time() - start)
    return {
        "success": True,
        "content_version_id": cv_id,
        "content_document_id": content_doc_id,
        "total_pages": total,
        "page_counts": page_counts,
    }


# ── /update-docx-tags — inject OX/OY offsets into .docx signing tags ──

MAX_DOCX_BASE64_BYTES = 70_000_000

# Regex for WSM bracket tags: [[TagKey]] or [[TagKey.suffix]]
# Captures: group(1)=tag key, group(2)=optional suffix (without leading dot)
_WSM_TAG_RE = re.compile(
    r'\[\[\s*([A-Za-z_]\w*(?:__[rc])?(?:\d+)?)'  # tag key (e.g. S1, mothermaiden_1, Field__c)
    r'(?:\.([^\]\s]*?))?'                          # optional .suffix
    r'\s*\]\]'
)

# Regex for DocuSign signing tags: \s1\, \i2\, etc.
_DS_SIGN_TAG_RE = re.compile(r'\\([sidntSIDNT])(\d+)\\')

# Regex for DocuSign text tabs: \fieldName_1_text\
_DS_TEXT_TAG_RE = re.compile(r'\\([a-zA-Z][a-zA-Z0-9]*)_(\d+)_(?:text|number|date|checkbox)\\', re.IGNORECASE)

# Regex to strip existing OX/OY from a suffix string
_STRIP_OXOY_RE = re.compile(r'OX-?\d+|OY-?\d+', re.IGNORECASE)


def _build_oxoy_suffix(ox: int, oy: int) -> str:
    """Build the OXnOYn suffix string."""
    return f"OX{ox}OY{oy}"


def _apply_offset_to_suffix(suffix: str, ox: int, oy: int) -> str:
    """Remove any existing OX/OY from suffix and append new values."""
    cleaned = _STRIP_OXOY_RE.sub("", suffix).strip(".")
    oxoy = _build_oxoy_suffix(ox, oy)
    if cleaned:
        return f"{cleaned}{oxoy}"
    return oxoy


def _update_tags_in_xml(xml_content: str, tag_offsets: dict) -> tuple[str, int]:
    """
    Process a single XML string, updating tags with OX/OY offsets.
    Returns (modified_xml, count_of_tags_updated).
    """
    count = 0

    # Build a lookup keyed by normalized tag key (case-insensitive for signing tags)
    # tag_offsets keys are already the tag keys (e.g. "S1", "mothermaiden_1")
    offsets_lower = {k.lower(): v for k, v in tag_offsets.items()}

    def replace_wsm_tag(m):
        nonlocal count
        tag_key = m.group(1)
        suffix = m.group(2) or ""

        # Look up by exact key first, then case-insensitive
        offsets = tag_offsets.get(tag_key) or offsets_lower.get(tag_key.lower())
        if not offsets:
            return m.group(0)  # no offset for this tag

        ox = offsets.get("ox", 0)
        oy = offsets.get("oy", 0)
        if ox == 0 and oy == 0:
            return m.group(0)  # skip zero offsets

        new_suffix = _apply_offset_to_suffix(suffix, ox, oy)
        count += 1
        return f"[[{tag_key}.{new_suffix}]]"

    def replace_ds_sign_tag(m):
        nonlocal count
        ds_letter = m.group(1)
        signer_num = m.group(2)
        tag_key = f"{ds_letter.upper()}{signer_num}"

        offsets = tag_offsets.get(tag_key) or offsets_lower.get(tag_key.lower())
        if not offsets:
            return m.group(0)

        ox = offsets.get("ox", 0)
        oy = offsets.get("oy", 0)
        if ox == 0 and oy == 0:
            return m.group(0)

        count += 1
        return f"[[{tag_key}.{_build_oxoy_suffix(ox, oy)}]]"

    def replace_ds_text_tag(m):
        nonlocal count
        field_name = m.group(1)
        signer_num = m.group(2)
        full_key = f"{field_name}_{signer_num}"

        offsets = tag_offsets.get(full_key) or offsets_lower.get(full_key.lower())
        if not offsets:
            # Also try just the field name
            offsets = tag_offsets.get(field_name) or offsets_lower.get(field_name.lower())
        if not offsets:
            return m.group(0)

        ox = offsets.get("ox", 0)
        oy = offsets.get("oy", 0)
        if ox == 0 and oy == 0:
            return m.group(0)

        count += 1
        return f"[[{full_key}.{_build_oxoy_suffix(ox, oy)}]]"

    # Apply replacements in order: DocuSign tags first (they get converted to WSM format),
    # then WSM tags (which may include freshly converted ones from a previous pass — but
    # since we do it in one pass per regex, order matters for non-overlapping matches).
    xml_content = _DS_TEXT_TAG_RE.sub(replace_ds_text_tag, xml_content)
    xml_content = _DS_SIGN_TAG_RE.sub(replace_ds_sign_tag, xml_content)
    xml_content = _WSM_TAG_RE.sub(replace_wsm_tag, xml_content)

    return xml_content, count


@app.post("/update-docx-tags")
async def update_docx_tags(request: Request):
    """
    Accept a .docx (as base64) and a map of tag offsets, then inject OX/OY
    position suffixes into all matching signing/text tags in the Word XML.
    """
    _verify_api_key(request)
    body = await request.json()
    start = time.time()

    import zipfile

    docx_b64 = body.get("docx_base64")
    tag_offsets = body.get("tag_offsets", {})

    if not docx_b64:
        raise HTTPException(status_code=400, detail="Missing docx_base64")
    if len(docx_b64) > MAX_DOCX_BASE64_BYTES:
        raise HTTPException(
            status_code=400,
            detail=f"DOCX too large ({len(docx_b64)} bytes, max {MAX_DOCX_BASE64_BYTES})",
        )
    if not tag_offsets:
        raise HTTPException(status_code=400, detail="Missing or empty tag_offsets")

    try:
        docx_bytes = base64.b64decode(docx_b64)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Invalid base64: {str(e)}")

    # Read the .docx ZIP
    try:
        in_zip = zipfile.ZipFile(io.BytesIO(docx_bytes), "r")
    except zipfile.BadZipFile:
        raise HTTPException(status_code=400, detail="Invalid .docx file (not a valid ZIP)")

    total_updated = 0
    out_buffer = io.BytesIO()

    with zipfile.ZipFile(out_buffer, "w", zipfile.ZIP_DEFLATED) as out_zip:
        for entry in in_zip.namelist():
            data = in_zip.read(entry)

            # Only process XML files inside word/ directory
            if entry.startswith("word/") and entry.endswith(".xml"):
                xml_content = data.decode("utf-8", errors="ignore")
                modified, count = _update_tags_in_xml(xml_content, tag_offsets)
                total_updated += count
                out_zip.writestr(entry, modified.encode("utf-8"))
            else:
                out_zip.writestr(entry, data)

    in_zip.close()

    result_b64 = base64.b64encode(out_buffer.getvalue()).decode("utf-8")
    logger.info(
        "update-docx-tags: tags_updated=%d duration=%.2fs",
        total_updated,
        time.time() - start,
    )
    return {
        "docx_base64": result_b64,
        "tags_updated": total_updated,
        "success": True,
    }


@app.get("/health")
async def health():
    return {"status": "ok"}
